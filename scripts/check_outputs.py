import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


PLACEHOLDER_RE = re.compile(
    r"\[(?:Paste|Insert|Continue|Optional|Short title|Journal|Exact filename|Reviewer|Concise|Completed evidence|Main manuscript|Section|Sentence|Reply|Missing values|Evidence checked)[^\]]*\]"
    r"|\b(?:TBD|TODO|FIGX|FIG\.\s*XX|FIGURE\s*XX|PAGE\s*XX|LINE\s*XX|TO BE ADDED)\b",
    re.IGNORECASE,
)


def rgb(run):
    color = run.font.color.rgb
    return str(color).upper() if color is not None else None


def paragraphs(doc):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip():
                        yield paragraph


def inspect_docx(path):
    doc = Document(path)
    ps = list(paragraphs(doc))
    text = "\n".join(p.text for p in ps)
    colors = {}
    italic_runs = 0
    for p in ps:
        for run in p.runs:
            value = rgb(run) or "INHERITED"
            colors[value] = colors.get(value, 0) + 1
            if run.italic:
                italic_runs += 1
    return {"paragraphs": ps, "text": text, "colors": colors, "italic_runs": italic_runs}


def normalize(text):
    # Preserve scientific case, units, and punctuation; normalize layout only.
    return re.sub(r"\s+", " ", text).strip()


def revised_quote_blocks(ps):
    """Find standalone quoted revision blocks inside REPLY sections.

    This is a conservative heuristic, not a complete quotation parser. Inline
    quotations and alternate reply labels still require manual inspection.
    """
    blocks, incomplete = [], []
    active_reply = False
    parts = []
    start = None
    closing = None
    boundary = re.compile(r"^(?:Q\d+\s*:|(?:Reviewer|Referee)\s*#?\s*\d+|(?:Editor|Reviewer|Referee)[’']?s?\s+Comments\s*:|\d+[.)]\s)", re.I)
    for index, p in enumerate(ps, 1):
        text = p.text.strip()
        if text.startswith("REPLY:") or boundary.match(text):
            if parts:
                incomplete.append(start)
                parts = []
            active_reply = text.startswith("REPLY:")
        if not active_reply:
            continue
        if not parts:
            if not text.startswith(('“', '"')):
                continue
            closing = '”' if text.startswith('“') else '"'
            start = index
            text = text[1:]
        candidate = text.rstrip()
        # Accept a closing quotation followed by a terminal period.
        if candidate.endswith(closing + '.'):
            candidate = candidate[:-1]
        if candidate.endswith(closing):
            parts.append(candidate[:-1])
            blocks.append({"paragraph": start, "text": "\n".join(parts)})
            parts = []
        else:
            parts.append(text)
    if parts:
        incomplete.append(start)
    return blocks, incomplete


def main():
    parser = argparse.ArgumentParser(description="Structural checks for lab reviewer-response outputs")
    parser.add_argument("--response", type=Path)
    parser.add_argument("--revision-map", dest="revision_map", type=Path)
    parser.add_argument("--revised-manuscript", dest="revised", type=Path)
    parser.add_argument("--supplement", type=Path, action="append", default=[],
                        help="revised SI DOCX to search for quotes; may be repeated")
    args = parser.parse_args()

    if not any((args.response, args.revision_map, args.revised)):
        parser.error("provide at least one DOCX")
    if args.supplement and not args.revised:
        parser.error("--supplement requires --revised-manuscript")

    report = {"errors": [], "warnings": [], "files": {}}

    response_info = None
    if args.response:
        response_info = inspect_docx(args.response)
        report["files"]["response"] = {
            "path": str(args.response),
            "colors": response_info["colors"],
            "italic_runs": response_info["italic_runs"],
        }
        text = response_info["text"]
        if "Point-by-point response" not in text:
            report["errors"].append("response title is missing")
        if not re.search(r"\b(?:Reviewer|Referee)\s*#?\s*\d+", text, re.IGNORECASE):
            report["errors"].append("reviewer/referee heading is missing")
        if "REPLY:" not in text:
            report["errors"].append("REPLY blocks are missing")
        if response_info["colors"].get("0070C0", 0) == 0:
            report["errors"].append("response contains no lab-blue #0070C0 text")
        if response_info["italic_runs"] == 0:
            report["errors"].append("response contains no italic reviewer or quotation text")
        hits = sorted(set(m.group(0) for m in PLACEHOLDER_RE.finditer(text)))
        if hits:
            report["errors"].append(f"response placeholders remain: {hits}")

    if args.revision_map:
        info = inspect_docx(args.revision_map)
        report["files"]["revision_map"] = {"path": str(args.revision_map)}
        required = (
            "Baseline manuscript",
            "Reviewer source",
            "Evidence status",
            "Stable location",
            "Operation",
            "Formatting instructions",
            "Response linkage",
            "Verification",
        )
        for field in required:
            if field not in info["text"]:
                report["errors"].append(f"revision map field is missing: {field}")
        hits = sorted(set(m.group(0) for m in PLACEHOLDER_RE.finditer(info["text"])))
        if hits:
            report["errors"].append(f"revision map placeholders remain: {hits}")

    if args.revised:
        revised_info = inspect_docx(args.revised)
        report["files"]["revised_manuscript"] = {
            "path": str(args.revised),
            "colors": revised_info["colors"],
        }
        if revised_info["colors"].get("FF0000", 0) == 0:
            report["errors"].append("revised manuscript contains no pure-red #FF0000 text")
        hits = sorted(set(m.group(0) for m in PLACEHOLDER_RE.finditer(revised_info["text"])))
        if hits:
            report["errors"].append(f"revised-manuscript placeholders remain: {hits}")

        if response_info:
            destinations = {str(args.revised): normalize(revised_info["text"])}
            for path in args.supplement:
                destinations[str(path)] = normalize(inspect_docx(path)["text"])
            report["files"]["supplements"] = [str(p) for p in args.supplement]
            quoted, incomplete = revised_quote_blocks(response_info["paragraphs"])
            report["quote_checks"] = []
            for block in quoted:
                quote = normalize(block["text"])
                matches = [path for path, text in destinations.items() if quote and quote in text]
                report["quote_checks"].append({"paragraph": block["paragraph"], "matched_files": matches})
                if not matches:
                    report["warnings"].append(
                        f"quoted revision at response paragraph {block['paragraph']} was not found "
                        "verbatim in supplied revised files; inspect wording, citations, and destination: "
                        + quote[:160]
                    )
            for paragraph in incomplete:
                report["warnings"].append(f"unclosed revision quotation at response paragraph {paragraph}")
            report["warnings"].append(
                "quote search is heuristic: manually check inline quotations, alternate reply labels, "
                "destination accuracy, and any SI not supplied; script success is not submission clearance"
            )

    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if report["errors"] else 0


if __name__ == "__main__":
    sys.exit(main())
