from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"
BLUE = RGBColor(0x00, 0x70, 0xC0)
RED = RGBColor(0xFF, 0x00, 0x00)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=120, bottom=70, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=GRAY, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def configure_document(doc, landscape=False):
    section = doc.sections[0]
    section.page_width = Inches(8.2677)
    section.page_height = Inches(11.6929)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(12)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_run(paragraph, text, *, color=BLACK, bold=False, italic=False, size=12):
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    return run


def build_response_template():
    doc = Document()
    configure_document(doc)
    p = doc.add_paragraph()
    add_run(p, "Point-by-point response to the reviewers’ comments:", bold=True)

    note = doc.add_paragraph()
    add_run(note, "* The reviewers’ comments are in Italic. Author’s responses are in ", italic=True)
    add_run(note, "blue", color=BLUE, italic=True)
    add_run(note, ". Quoted revisions are in ", italic=True)
    add_run(note, "blue italics", color=BLUE, italic=True)
    add_run(note, ". All revisions are highlighted in ", italic=True)
    add_run(note, "red", color=RED, italic=True)
    add_run(note, " in the revised manuscript.", italic=True)

    p = doc.add_paragraph()
    add_run(p, "REVIEWER COMMENTS", bold=True)
    p = doc.add_paragraph()
    add_run(p, "Reviewer #1 (Comments to the Author):", bold=True)
    p = doc.add_paragraph()
    add_run(p, "[Paste the reviewer’s general assessment or numbered comment verbatim.]", italic=True)

    p = doc.add_paragraph()
    add_run(p, "REPLY:", color=BLUE, bold=True)
    add_run(p, " [Acknowledge the point briefly and answer it directly. State the experiment or analysis, result, interpretation, and limitation when applicable.]", color=BLUE)

    p = doc.add_paragraph()
    add_run(p, "In the revised manuscript, under “[section heading]”, paragraph [stable paragraph description], we added/revised the following text:", color=BLUE)
    p = doc.add_paragraph()
    add_run(p, "“[Insert the exact revised manuscript text here.]”", color=BLUE, italic=True, size=11)

    p = doc.add_paragraph()
    add_run(p, "[Optional response-only figure or table: Figure R1-1. State what it shows and identify its final destination in the revised manuscript.]", color=BLUE, size=11)

    doc.add_paragraph()
    p = doc.add_paragraph()
    add_run(p, "Reviewer #2 (Comments to the Author):", bold=True)
    p = doc.add_paragraph()
    add_run(p, "[Continue in the original reviewer order and repeat the same structure.]", italic=True)

    doc.save(ASSETS / "lab-response-template.docx")


def add_field_row(table, label, value, shade=False):
    cells = table.add_row().cells
    cells[0].width = Inches(1.8)
    cells[1].width = Inches(7.8)
    for cell in cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        set_cell_borders(cell)
        if shade:
            set_cell_shading(cell, "EAF2F8")
    p = cells[0].paragraphs[0]
    add_run(p, label, bold=True, size=9.5)
    p = cells[1].paragraphs[0]
    add_run(p, value, size=9.5)


def build_revision_map_template():
    doc = Document()
    configure_document(doc, landscape=True)
    p = doc.add_paragraph()
    add_run(p, "Manuscript Revision Map", bold=True)
    p = doc.add_paragraph()
    add_run(p, "Purpose: ", bold=True)
    add_run(p, "Translate reviewer responses and new evidence into precise manuscript edits. Apply this map only to the baseline manuscript identified below.")

    metadata = doc.add_table(rows=0, cols=2)
    metadata.alignment = WD_TABLE_ALIGNMENT.CENTER
    metadata.autofit = False
    add_field_row(metadata, "Manuscript short title", "[Short title]", shade=True)
    add_field_row(metadata, "Journal and round", "[Journal; revision round]", shade=True)
    add_field_row(metadata, "Baseline manuscript", "[Exact filename of the originally submitted manuscript]", shade=True)
    add_field_row(metadata, "Evidence inputs", "[Reviewer-comments file; new-data files; supervisor instructions]", shade=True)
    add_field_row(metadata, "Map status", "[Draft / author-confirmed / applied; date]", shade=True)

    doc.add_page_break()
    p = doc.add_paragraph()
    add_run(p, "Global manuscript formatting rules", bold=True)
    for text in (
        "Color only inserted or replacement text pure red #FF0000.",
        "Keep unchanged text, styles, fields, figures, tables, and page setup intact.",
        "Record deletions explicitly in this map; use true Track Changes only if requested.",
        "Preserve live citation and bibliography fields. Do not invent missing references.",
        "Synchronize every quoted revision in the response with the final manuscript wording.",
    ):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        add_run(p, "• ", bold=True)
        add_run(p, text)

    doc.add_paragraph()
    p = doc.add_paragraph()
    add_run(p, "Change R1.1-M1", bold=True)
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    fields = [
        ("Reviewer source", "[Reviewer 1, Comment 1]"),
        ("Reviewer concern", "[Concise issue statement]"),
        ("Evidence status", "[Completed evidence / planned work / proposed wording / supervisor instruction / unresolved]"),
        ("Evidence basis", "[Exact file, dataset, sample, condition, value, figure or table]"),
        ("Target document", "[Main manuscript / Supplementary Information / figure / table / source data]"),
        ("Stable location", "[Section > subsection > paragraph purpose > exact anchor phrase > insertion point]"),
        ("Operation", "[Insert / replace / delete / move / add figure or table / revise legend / add method / add citation]"),
        ("Change form", "[Sentence / paragraph / methods subsection / figure panel / table row / legend / limitation statement]"),
        ("Exact text or content specification", "[Exact proposed wording, or a bounded specification pending data confirmation]"),
        ("Figure or table instructions", "[Response-only number, final number, panel, legend, statistics, source data and callout]"),
        ("Formatting instructions", "[Changed text red #FF0000; inherit local style; preserve citation fields; any special requirement]"),
        ("Response linkage", "[Reply paragraph and exact quoted revision that must match]"),
        ("Dependencies or confirmation", "[Missing values, final numbering, citation, statistical result or supervisor decision]"),
        ("Verification", "[Evidence checked; inserted; response synchronized; fields preserved; visual QA passed]"),
    ]
    for idx, (label, value) in enumerate(fields):
        add_field_row(table, label, value, shade=idx % 2 == 0)

    doc.save(ASSETS / "manuscript-revision-map-template.docx")


if __name__ == "__main__":
    ASSETS.mkdir(parents=True, exist_ok=True)
    build_response_template()
    build_revision_map_template()
    print(f"Created templates in {ASSETS}")
