"""Synthetic quote-parser regressions; no manuscript data is included."""
import sys
import json
import subprocess
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.shared import RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from check_outputs import normalize, revised_quote_blocks


def parse(*texts):
    return revised_quote_blocks([SimpleNamespace(text=t) for t in texts])


class QuotationTests(unittest.TestCase):
    def test_reviewer_quote_is_not_an_author_revision(self):
        blocks, pending = parse('Reviewer #1:', '“Please explain this statement.”',
                                'REPLY: Clarified.', '“The outcome was measured directly.”')
        self.assertEqual([b['text'] for b in blocks], ['The outcome was measured directly.'])
        self.assertEqual(pending, [])

    def test_multiline_quote_includes_later_paragraphs(self):
        blocks, pending = parse('REPLY: The revised methods read:',
                                '“First preparation step.', 'Second preparation step.”')
        self.assertEqual(blocks[0]['text'], 'First preparation step.\nSecond preparation step.')
        self.assertEqual(pending, [])
        self.assertNotIn(normalize(blocks[0]['text']), normalize('First preparation step.'))

    def test_unclosed_quote_stops_at_next_question(self):
        blocks, pending = parse('REPLY: Revised:', '“Unclosed passage.',
                                'Q2: Another concern.', 'REPLY: Fixed.', '“New wording.”')
        self.assertEqual(pending, [2])
        self.assertEqual([b['text'] for b in blocks], ['New wording.'])

    def test_incomplete_quote_at_end_is_reported(self):
        self.assertEqual(parse('REPLY: Revised:', '“Unfinished.')[1], [2])

    def test_straight_quotes_and_terminal_period(self):
        blocks, pending = parse('REPLY: Revised:', '"A short passage".')
        self.assertEqual(blocks[0]['text'], 'A short passage')
        self.assertEqual(pending, [])

    def test_scientific_case_is_preserved(self):
        self.assertNotEqual(normalize('5 mM'), normalize('5 mm'))
        self.assertEqual(normalize('A  paragraph.\nNext sentence.'), 'A paragraph. Next sentence.')

    def test_each_reviewers_comments_reset_state(self):
        blocks, pending = parse('REPLY: No change.', 'Reviewer #2 Comments:',
                                '“This is a reviewer quotation.”', 'REPLY: Revised:', '“New text.”')
        self.assertEqual([b['text'] for b in blocks], ['New text.'])
        self.assertEqual(pending, [])


class DestinationTests(unittest.TestCase):
    def test_cli_reports_si_destination_and_missing_destination(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            response, main, si = [root / name for name in ('response.docx', 'main.docx', 'si.docx')]
            doc = Document()
            doc.add_paragraph('Point-by-point response')
            doc.add_paragraph('Reviewer #1:')
            doc.add_paragraph('Q1: Describe the measurement.').runs[0].italic = True
            doc.add_paragraph('REPLY: Added to the supplement.').runs[0].font.color.rgb = RGBColor.from_string('0070C0')
            doc.add_paragraph('“The measurement used a reference standard.”')
            doc.save(response)
            doc = Document()
            doc.add_paragraph('The main text was revised.').runs[0].font.color.rgb = RGBColor.from_string('FF0000')
            doc.save(main)
            doc = Document()
            doc.add_paragraph('The measurement used a reference standard.')
            doc.save(si)
            command = [sys.executable, str(Path(__file__).resolve().parents[1] / 'scripts' / 'check_outputs.py'),
                       '--response', str(response), '--revised-manuscript', str(main)]
            missing = subprocess.run(command, capture_output=True, text=True, encoding='utf-8', check=True)
            report = json.loads(missing.stdout)
            self.assertEqual(report['quote_checks'][0]['matched_files'], [])
            self.assertTrue(any('not found verbatim' in w for w in report['warnings']))
            supplied = subprocess.run(command + ['--supplement', str(si)], capture_output=True,
                                      text=True, encoding='utf-8', check=True)
            report = json.loads(supplied.stdout)
            self.assertEqual(report['quote_checks'][0]['matched_files'], [str(si)])
            self.assertFalse(any('not found verbatim' in w for w in report['warnings']))


if __name__ == '__main__':
    unittest.main()
